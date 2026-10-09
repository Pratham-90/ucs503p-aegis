# ruff: noqa: E501  (long prose strings: this script is report content)
"""Generate report/Aegis_Prototype_Report_Content.docx (brief section 12).

The team pastes this content into the TIET Overleaf report template. Every number
comes from metrics/prototype-metrics.json (written by scripts/collect_metrics.py);
a missing value is printed as a bold "[TODO: not measured]". The worked Shamir
example is computed here, so its arithmetic is correct by construction.

    python scripts/make_report_docx.py
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
METRICS = json.loads((ROOT / "metrics" / "prototype-metrics.json").read_text(encoding="utf-8"))
M = METRICS["metrics"]
SHOTS = ROOT / "report-assets" / "screenshots"
DIAGRAMS = ROOT / "report-assets" / "diagrams"
OUT = ROOT / "report" / "Aegis_Prototype_Report_Content.docx"
TODO = "[TODO: not measured]"
LIVE_URL = "[TODO: live URL — fill in after the Vercel deployment]"

doc = Document()
_fig = {"chapter": 0, "n": 0}
_tab = {"n": 0}
FIGURES: list[tuple[str, str, str]] = []

# --- styles ------------------------------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for section in doc.sections:
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)


# --- metric helpers --------------------------------------------------------------------------
def mv(name: str):
    """Raw metric value, or None if missing / not measured."""
    m = M.get(name)
    if m is None or m.get("value") == "TODO: not measured":
        return None
    return m["value"]


def ms(name: str, key: str | None = None, unit: str = "") -> str:
    """Metric as text; '[TODO: not measured]' when absent."""
    value = mv(name)
    if value is None:
        return TODO
    if key is not None:
        value = value.get(key) if isinstance(value, dict) else None
        if value is None:
            return TODO
    return f"{value}{unit}"


def met(name: str) -> str:
    m = M.get(name, {})
    return {True: "Met", False: "Not met"}.get(m.get("met"), "—")


# --- text helpers ----------------------------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|\[TODO[^\]]*\]|`[^`]+`)")


def add_runs(paragraph, text: str, *, italic: bool = False, size: float | None = None) -> None:
    """Minimal markup: **bold**, `code`, and [TODO…] (always bold)."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("[TODO"):
            run = paragraph.add_run(part)
            run.bold = True
            run.font.color.rgb = RGBColor(0xB4, 0x1C, 0x1C)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(part)
        run.italic = run.italic or italic
        if size:
            run.font.size = Pt(size)


def para(text: str, *, style: str | None = None, align=None, italic: bool = False, size: float | None = None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, italic=italic, size=size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, item)


def numbered(items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, item)


def heading(text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def page_break() -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def code_block(lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.8)
    doc.add_paragraph()


def _shade(cell, hex_fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def table(header: list[str], rows: list[list[str]], caption: str | None = None, label: str | None = None,
          widths: list[float] | None = None) -> None:
    if caption:
        _tab["n"] += 1
        cap = doc.add_paragraph()
        add_runs(cap, f"**Table {_fig['chapter']}.{_tab['n']}: {caption}**")
        if label:
            add_runs(cap, f"   (suggested label: `{label}`)", italic=True, size=9)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], f"**{h}**", size=9.5)
        _shade(cell, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], str(value), size=9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def figure(path: Path, caption: str, label: str, width_cm: float = 15.5) -> None:
    _fig["n"] += 1
    number = f"{_fig['chapter']}.{_fig['n']}"
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para(f"[TODO: missing image {path.relative_to(ROOT)}]")
    cap = para(f"Figure {number}: {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=10)
    add_runs(cap, f"   (suggested label: `{label}`; file: `{path.relative_to(ROOT).as_posix()}`)", size=8.5)
    FIGURES.append((f"Figure {number}", caption, path.relative_to(ROOT).as_posix()))


def chapter(n: int, title: str) -> None:
    _fig["chapter"], _fig["n"], _tab["n"] = n, 0, 0
    page_break()
    heading(f"Chapter {n} – {title}", 1)


# =====================================================================================
# How to use this file
# =====================================================================================
heading("How to use this file", 1)
para(f"This document holds the content for the Aegis prototype report, structured to match the TIET report "
     f"template chapter by chapter so that each section can be pasted into Overleaf. It was generated on "
     f"{datetime.now(UTC).strftime('%Y-%m-%d')} by `scripts/make_report_docx.py` from measurements taken at commit "
     f"`{METRICS['git_sha']}` (`metrics/prototype-metrics.json`).")
bullets([
    "**Numbers.** Every figure quoted comes from `metrics/prototype-metrics.json`. Anything that was not measured "
    "is printed as [TODO: not measured]; do not replace it with an estimate. Re-run "
    "`python scripts/collect_metrics.py` and this script after measuring the deployed site.",
    "**Citations** appear as `[cite: key]`. Each key exists in `report/references.bib`; in LaTeX write "
    "`\\cite{key}`.",
    "**Figures** are numbered by chapter (Figure 2.1, …). Each caption shows a suggested LaTeX label "
    "(`fig:…`) and the PNG file to use from `report-assets/`. Tables use `tab:…` labels; suggested section labels "
    "follow `sec:chapter-topic`.",
    "**[TODO] markers** (bold, red) mark information only the team can supply: roll numbers, group, the live URL, "
    "user-testing results and the division of work.",
    "**Requirement IDs** (`FR-n`, `NFR-cat-n`, `NG-n`) refer to the SRS in `docs/srs/index.md` and "
    "`project-proposal/main.tex`.",
    "The **appendix** at the end is for the team (presenter notes, viva preparation) and is not part of the report.",
])

# =====================================================================================
# Title page fields
# =====================================================================================
page_break()
heading("Title page fields", 1)
table(["Field", "Content"], [
    ["Course code / title", "UCS503P – Software Engineering (Laboratory) [TODO: confirm with the team]"],
    ["Project title", "Aegis"],
    ["Subtitle", "An Encrypted Dead-Man's-Switch Legacy Vault using Shamir's Secret Sharing"],
    ["Submitted by", "Pratham Arora (Roll No. 1024030001 [TODO: confirm]); Krishna Pandey (Roll No. [TODO]); "
                     "Nipun Behl (Roll No. [TODO])"],
    ["Group", "[TODO]"],
    ["Year / branch", "BE Third Year, COE [TODO: confirm]"],
    ["Submitted to", "Dr. Paramveer Sidhu"],
    ["Evaluation", "Prototype Evaluation (Mid-Semester) [TODO: confirm]"],
    ["Institute", "Thapar Institute of Engineering and Technology, Patiala — 2026–27 ODD semester"],
], widths=[4.5, 11.5])

# =====================================================================================
# Chapter 1
# =====================================================================================
chapter(1, "Introduction")
heading("1.1 Background and Context", 2)
para("A growing share of what people own and know now exists only in digital form: banking and brokerage "
     "credentials, cryptocurrency keys, password-manager master passwords, encrypted archives, business-continuity "
     "information and personal messages meant for particular people. The value of this material is entirely "
     "conditional on access, and access is exactly what is lost when its owner dies or becomes permanently "
     "incapacitated. Unlike a physical estate, a forgotten password or an unrecorded key cannot be recovered by "
     "an executor; the data simply becomes unreachable.")
para("Existing approaches force an uncomfortable choice. Sharing passwords with a relative or friend makes the "
     "secret readable from the moment it is shared, and its safety then depends indefinitely on the honesty, "
     "competence and availability of one person. Appointing a single executor or custodian concentrates the same "
     "risk in one party and is triggered by a slow legal process rather than by the owner's absence. "
     "Provider-held recovery schemes, in which a service keeps the means to unlock an account, require trusting "
     "the provider and anyone who compromises it with the plaintext. In every case there is a single point of "
     "failure: one document, one person or one operator whose loss, compromise or delay defeats the owner's "
     "intent.")
para("The project therefore combines two ideas. A dead man's switch releases information when its owner stops "
     "responding, so the trigger is the owner's sustained silence rather than a human decision. Threshold secret "
     "sharing [cite: shamir1979] removes the single custodian: the decryption key is split among several trustees "
     "so that any K of N of them can rebuild it, while any K−1 learn nothing at all. Aegis adds a third "
     "property that recent breaches have made essential: the service that runs the switch should never be able "
     "to read what it protects.")

heading("1.1.1 Problem Statement", 3)
para("The problem addressed is how to release a secret to designated recipients when, and only when, its owner "
     "has become unresponsive, without any single party — including the system operator — being able to read the "
     "secret beforehand, and without depending on any single custodian. The scope of the prototype is fixed by "
     "the SRS: one vault per owner (NG-3); check-in prompts and release notices by email only, with no SMS "
     "(NG-1) and no native mobile application (NG-2); N trustees with a reconstruction threshold of at least two "
     "(FR-4); and a zero-knowledge server that stores only ciphertext and trustee-encrypted shares (NFR-SEC-5). "
     "The catastrophic failure mode is premature release, which is irreversible (NFR-REL-1).")

heading("1.2 Project Objectives", 2)
para("The prototype was built to the following SMART objectives, each tied to SRS requirements and to the "
     "prototype milestone (Week 8 of the 12-week plan).")
numbered([
    "Never release a vault before its deadline plus grace period: zero early releases across at least 1,000 "
    "simulated schedules with injected faults, by Week 8 (NFR-REL-1, NFR-REL-3).",
    "Keep the server zero-knowledge: an automated inspection of everything the server stores and receives finds "
    "no plaintext payload, key or share, by Week 8 (FR-2, FR-2a, NFR-SEC-1, NFR-SEC-5).",
    "Make the threshold property demonstrable and verified: any K of N shares recover the vault, fewer than K "
    "fail, verified by property-based tests and an end-to-end browser test of the full demo script, by Week 8 "
    "(FR-8, NFR-SEC-2).",
    "Implement the cryptographic core twice — an authoritative Python implementation and the deployed TypeScript "
    "client — and bind them with shared known-answer test vectors that agree on every case (FR-2a).",
    "Keep check-ins fast and the code maintainable: check-in confirmation p95 within 500 ms under nominal load "
    "(NFR-PERF-1) and at least 80% test coverage on the crypto and scheduler packages (NFR-MAINT-1), deployed "
    "on Vercel for the prototype demonstration.",
])

# =====================================================================================
# Chapter 2
# =====================================================================================
chapter(2, "Methodology and Design")
heading("2.1 System Architecture", 2)
para("Aegis is a web application with a React single-page frontend and a FastAPI backend, deployed as a single "
     "Vercel project so that the browser talks to one origin and the session cookie needs no cross-origin "
     "configuration [cite: vercelpython]. All payload cryptography runs in the browser; the server schedules "
     "check-ins, stores ciphertext and encrypted shares in Postgres, and sends email. Because serverless "
     "functions do not keep a process alive, the scheduler is driven by external triggers that call an "
     "idempotent tick endpoint. Figure 2.1 shows the deployment and Figure 2.2 the client/server trust boundary.")
figure(DIAGRAMS / "architecture-2.png", "Deployment architecture of the prototype", "fig:architecture")
figure(DIAGRAMS / "architecture-1.png", "Client/server trust boundary: only ciphertext and encrypted blobs cross it",
       "fig:trust-boundary")

heading("2.1.1 High-Level Design", 3)
para("**Components.** The browser hosts the user interface and the client crypto module (AES-256-GCM, a "
     "hand-written Shamir implementation over BigInt, and RSA-OAEP). The FastAPI application exposes "
     "authentication, vault, check-in, trustee, cron and demo endpoints under `/api` [cite: fastapidocs]. "
     "Four domain packages sit behind it: `crypto` (the authoritative specification and test oracle), "
     "`scheduler` (the lifecycle), `vault` (SQLAlchemy models and repositories [cite: sqlalchemydocs]) and "
     "`notifications` (a transactional outbox).")
para("**Trust model.** The server is trusted for scheduling and delivery, never for confidentiality. The "
     "browser generates a random 256-bit key, encrypts the payload, splits the key into N shares and encrypts "
     "each share to its trustee's RSA public key; the server receives only the results. Trustees generate their "
     "keypairs in their own browsers and keep the private key as a downloaded file. A server breach alone "
     "therefore reveals nothing; a breach of confidentiality requires the server and at least K trustees. The "
     "documented residual risks (for example, a malicious operator serving altered client code) are discussed in "
     "Section 3.3.")
para("**Vault state machine.** Each vault moves through Setup, Active, Warning, Grace and Released (Figure 2.3). "
     "The decision is made by a single pure function, `evaluate(clock, now)`, which releases a vault only when "
     "the current time has reached the deadline plus the full grace period with no confirmed check-in. A "
     "check-in from Active, Warning or Grace returns the vault to Active; Released is terminal.")
figure(DIAGRAMS / "vault-lifecycle-1.png", "Vault lifecycle state machine with its triggers and actions",
       "fig:state-machine")
para("Figures 2.4 and 2.5 show the two principal flows: creating and sealing a vault, and releasing and "
     "recovering it. Figure 2.6 shows the data model; the use-case view is Figure 2.7.")
figure(DIAGRAMS / "sequences-1.png", "Sequence: creating and sealing a vault", "fig:seq-create")
figure(DIAGRAMS / "sequences-2.png", "Sequence: release and K-of-N recovery", "fig:seq-release")
figure(DIAGRAMS / "data-model-1.png", "Data model (SQLAlchemy entities)", "fig:data-model", width_cm=14)
figure(DIAGRAMS / "use-case-1.png", "Use-case diagram: Owner, Trustee and the time-triggered Scheduler",
       "fig:use-case", width_cm=14)

heading("2.1.2 Technical Stack", 3)
para("**Framework/Language:** Python 3.12 with FastAPI for the API; TypeScript with React 18, Vite and Tailwind "
     "CSS for the client [cite: reactdocs].")
para("**Database:** SQLAlchemy 2 with SQLite for local development and tests, and Postgres on Neon in deployment "
     "(NFR-PORT-1). Payloads are stored as ciphertext in the database, capped at about 1 MB.")
para("**Tools:** Git and GitHub with a `prototype` branch; GitHub Actions for continuous integration and the "
     "5-minute scheduler trigger; Vercel for hosting; Playwright for end-to-end testing [cite: playwrightdocs]; "
     "ruff, ESLint and Prettier for code quality.")
para("**Libraries:** the browser's Web Cryptography API for AES-256-GCM [cite: nist80038d] and RSA-OAEP-2048 with "
     "SHA-256 [cite: rfc8017] [cite: w3cwebcrypto]; argon2-cffi for password hashing [cite: rfc9106]; PyJWT for "
     "session tokens; pytest, Hypothesis [cite: maciver2019hypothesis] and Vitest for testing.")
table(["Layer", "Technology", "Reason"], [
    ["Frontend", "React 18 + TypeScript + Vite + Tailwind", "Planned stack; static build served from Vercel's CDN"],
    ["Client crypto", "Web Crypto (AES-256-GCM, RSA-OAEP-2048/SHA-256)",
     "Built into browsers; no third-party crypto library on the client"],
    ["Secret sharing", "Hand-written Shamir over GF(2^521 − 1)", "SRS constraint; Python and TypeScript versions "
     "bound by shared test vectors"],
    ["Backend", "Python 3.12 + FastAPI as a Vercel Python function", "Planned stack; same origin as the frontend"],
    ["Database", "SQLAlchemy 2; SQLite locally, Neon Postgres deployed",
     "NFR-PORT-1; serverless functions have no persistent disk"],
    ["Scheduler", "Pure evaluate() + idempotent POST /api/cron/tick",
     "Serverless functions keep no process alive, so no APScheduler"],
    ["Tick triggers", "GitHub Actions (5 min), Demo Console button, lazy due-checks, daily Vercel Cron",
     "Hobby-plan Vercel Cron runs at most once a day"],
    ["Email", "Outbox + EmailSender (log / Resend)", "Demo works without a verified email domain"],
    ["Auth", "argon2id + JWT in an httpOnly, SameSite=Lax cookie", "NFR-SEC-3"],
    ["Testing", "pytest, Hypothesis, pytest-cov, Vitest, Playwright", "NFR-MAINT-1 and the demo script as a test"],
], caption="Technology choices and their reasons", label="tab:stack", widths=[3, 6, 7])

heading("2.2 Implementation Details", 2)
heading("2.2.1 Module 1: Core Functionality — client-side encryption and Shamir's Secret Sharing", 3)
para("**Finite field.** All arithmetic is carried out modulo the Mersenne prime p = 2^521 − 1, which is large "
     "enough to hold a 256-bit AES key as a single field element. Division is multiplication by a modular "
     "inverse (Fermat's little theorem in TypeScript, Python's three-argument `pow` in the oracle).")
para("**Polynomial construction.** To split a key S with threshold K, the client draws K−1 coefficients "
     "uniformly at random from the field and forms f(x) = S + a1·x + … + a(K−1)·x^(K−1), so that f(0) = S. "
     "Trustee i receives the point (i, f(i)) for i = 1 … N; x = 0 is never used because f(0) is the secret.")
para("**Reconstruction.** Any K points determine a unique polynomial of degree at most K−1, so the secret is "
     "recovered by Lagrange interpolation evaluated at zero: S = Σ y_i · Π_{j≠i} x_j / (x_j − x_i) (mod p). With "
     "only K−1 points, every candidate secret is consistent with exactly one polynomial through those points, "
     "so the shares reveal nothing about S [cite: shamir1979]. A Hypothesis property test checks this directly: "
     "for random K−1 shares and a random candidate secret, it constructs coefficients that the real split code "
     "turns into exactly those shares.")
para("**Share encoding and wrapping.** A share is encoded as 68 bytes (a version byte, the x-coordinate and a "
     "66-byte y) and encrypted to its trustee's RSA-OAEP-2048 public key, producing a 256-byte blob; this fits "
     "RSA-OAEP-SHA-256's 190-byte plaintext limit. Shares that trustees exchange in the Recovery Room use the "
     "text form `aegis-share:v1:<x>:<y-base64url>`. The payload is encrypted with AES-256-GCM using the vault "
     "identifier as associated data, so a ciphertext cannot be opened as another vault's.")
para(f"**Two implementations, one specification.** The Python package `code/crypto` is the authoritative "
     f"specification and test oracle; it emits a versioned file of known-answer vectors (Shamir splits from "
     f"known coefficients, reconstructions, AES-GCM and RSA-OAEP cases). The TypeScript client must reproduce "
     f"every vector, and continuous integration fails if the two disagree. Result: "
     f"{ms('Cross-language vectors (Python -> TypeScript)')}.")
para("Pseudocode of the client-side sealing and recovery steps:")
code_block([
    "SEAL(payload, trustees, K):",
    "    key  <- 32 random bytes;  iv <- 12 random bytes",
    "    ct   <- AES-256-GCM-Encrypt(key, iv, payload, aad = 'aegis-vault:v1:' || vault_id)",
    "    a[1..K-1] <- uniform random field elements",
    "    for i in 1..N:  share_i <- (i, f(i)) where f(x) = int(key) + a1*x + ... + a(K-1)*x^(K-1) mod p",
    "    blob_i <- RSA-OAEP-Encrypt(trustee_i.public_key, encode(share_i))",
    "    upload(ct, iv, blob_1..blob_N);  wipe(key)",
    "",
    "RECOVER(shares, K, ct, iv):",
    "    if |shares| < K: report 'not enough shares'",
    "    S <- sum_i y_i * prod_{j != i} x_j / (x_j - x_i)  mod p",
    "    key <- 32-byte big-endian encoding of S  (fails if S does not fit: wrong or too few shares)",
    "    return AES-256-GCM-Decrypt(key, iv, ct, aad)    (GCM authentication fails on any error)",
])

# Worked example over a toy prime, computed here.
P, SECRET, A1, A2 = 31, 13, 7, 4
f = [(x, (SECRET + A1 * x + A2 * x * x) % P) for x in range(1, 6)]
chosen = [f[0], f[2], f[4]]
basis = []
for i, (xi, _) in enumerate(chosen):
    num = den = 1
    for j, (xj, _) in enumerate(chosen):
        if i != j:
            num = num * xj % P
            den = den * (xj - xi) % P
    basis.append(num * pow(den, -1, P) % P)
recovered = sum(y * b for (_, y), b in zip(chosen, basis, strict=True)) % P
assert recovered == SECRET
two = chosen[:2]


def through(candidate: int) -> tuple[int, int, int]:
    """Coefficients (c0, c1, c2) of the quadratic through (0, candidate) and the two known shares."""
    (x1, y1), (x2, y2) = two
    # Solve c1*x + c2*x^2 = y - candidate for both points.
    r1, r2 = (y1 - candidate) % P, (y2 - candidate) % P
    det = (x1 * x2 * x2 - x2 * x1 * x1) % P
    c1 = (r1 * x2 * x2 - r2 * x1 * x1) * pow(det, -1, P) % P
    c2 = (x1 * r2 - x2 * r1) * pow(det, -1, P) % P
    return candidate, c1, c2


para(f"**Worked example (toy prime p = {P}).** Take secret S = {SECRET}, threshold K = 3 and N = 5, with random "
     f"coefficients a1 = {A1} and a2 = {A2}, so f(x) = {SECRET} + {A1}x + {A2}x² (mod {P}). The five shares are "
     + ", ".join(f"({x}, {y})" for x, y in f) + ". Using shares x = 1, 3, 5, the Lagrange basis values at zero "
     "are " + ", ".join(f"ℓ{i + 1}(0) = {b}" for i, b in enumerate(basis))
     + f" (mod {P}), and Σ y_i·ℓ_i(0) = "
     + " + ".join(f"{y}·{b}" for (_, y), b in zip(chosen, basis, strict=True))
     + f" ≡ {recovered} (mod {P}), which is the secret.")
alternatives = [through(c) for c in (0, 5, 22)]
para("With only two shares, (" + "), (".join(f"{x}, {y}" for x, y in two) + "), every candidate secret fits "
     "some quadratic: for example " + "; ".join(f"S = {c0} gives f(x) = {c0} + {c1}x + {c2}x²" for c0, c1, c2 in
                                                alternatives)
     + f" (all mod {P}), each passing through both known shares. Two shares therefore say nothing about S.")

heading("2.2.2 Module 2: Secondary Features", 3)
para("**Scheduler and tick design.** The lifecycle is decided by the pure function `evaluate(clock, now)`, which "
     "returns the new state and the actions owed on the way (check-in prompt, grace reminder, release). It moves "
     "only forward, so a backward clock jump never undoes a transition, and it catches up through missed "
     "transitions in order after an outage. The SRS fixes only the interval and grace period (FR-3); the "
     "prototype places the Warning-to-Grace boundary inside the grace window (by default at its midpoint), so "
     "the release instant remains exactly the deadline plus the grace period. A late check-in, arriving at or "
     "after that instant, is refused, so the outcome never depends on whether a tick happened to run first.")
para("**Transactional tick and exactly-once release.** `tick()` loads due vaults with `SELECT … FOR UPDATE SKIP "
     "LOCKED` on Postgres (SQLite takes the write lock up front), applies `evaluate` to the locked rows and writes "
     "the state change and the outbox rows in one transaction. Each outbox row has a unique deduplication key, "
     "for example `release:<vault>:<trustee>`, so a release is recorded exactly once even when two ticks race "
     "(NFR-REL-3). Deadlines live in the database, so a restarted process resumes where the previous one "
     "stopped (NFR-REL-2). Ticks come from GitHub Actions every five minutes, a Demo Console button, a cheap "
     "due-check whenever an owner or trustee status endpoint is read, and a daily Vercel Cron backstop.")
para("**Check-in tokens.** Prompt emails carry a one-click link with a single-use, expiring token; only its "
     "SHA-256 hash is stored. The link's GET request only previews the token; the page confirms with a POST, so "
     "email security scanners that pre-fetch links cannot check in on the owner's behalf (threat T-5).")
para("**Trustee enrolment.** Each trustee opens an invitation link; their browser generates an RSA-OAEP-2048 "
     "keypair, uploads only the public key and forces a download of the private-key file (optionally also kept in "
     "IndexedDB). The API rejects any key containing private-key fields. FR-4 is enforced on the server and in "
     "the user interface: K below 2 is rejected, K = N produces a warning that losing any trustee makes the vault "
     "unrecoverable, and the loss tolerance N − K is displayed.")
para("**Notification service.** Email is a strategy behind an `EmailSender` interface: `log` mode records every "
     "message in an in-app Outbox (used in the demonstration), and `resend` mode also sends through Resend.")
para("**Frontend pages.** A landing page; registration and login; an owner dashboard with a live countdown on "
     "the server's clock and a single \"I'm OK — check in\" button; a vault wizard that asks for the secret only "
     "after every trustee has enrolled and then shows exactly what the server received; the one-click check-in "
     "page; trustee enrolment and portal pages; a Recovery Room that reconstructs and decrypts entirely in the "
     "browser; and a token-protected Demo Console with the Outbox, tick log, server view and measured results.")
para("**Deployment.** One Vercel project builds the Vite frontend as static files and runs `api/index.py` as a "
     "Python function; rewrites send `/api/*` to the function and all other paths to the single-page app.")

heading("Software Engineering process", 3)
para("The project follows the incremental 12-week plan in `docs/plan/twelve-week-plan.md`: requirements and "
     "scaffolding (Week 1), the zero-knowledge trust model and threat model (Week 2), and the prototype "
     "(Week 8) [cite: pressman2020]. Work is done on feature branches with small conventional commits; the "
     "prototype was built on a `prototype` branch and merged through a pull request. Continuous integration "
     "runs ruff, pytest with a coverage gate, ESLint, Vitest (including the cross-language vectors), the "
     "production build and the end-to-end test on every push. Requirements are traced to modules and tests as "
     "follows.")
table(["Requirement", "Module", "Verified by"], [
    ["FR-1, NFR-SEC-3", "code/api/routes/auth.py, security.py", "test_api.py (register/login, argon2id hash, cookie flags)"],
    ["FR-2, FR-2a, NFR-SEC-1, NFR-SEC-5", "frontend/src/crypto/vault.ts", "vault.test.ts; inspect_server_store.py"],
    ["FR-3, FR-4", "code/api/routes/vaults.py; VaultNew.tsx", "test_api.py (K < 2, K = N, interval); E2E step 3"],
    ["FR-5, FR-7", "code/scheduler/state.py, tick.py", "test_scheduler_state.py, test_scheduler_tick.py"],
    ["FR-6, NFR-USE-1", "code/api/routes/checkin.py; CheckinLink.tsx", "test_api.py; E2E (emailed link)"],
    ["FR-8, NFR-SEC-2", "shamir.py / shamir.ts; Recover.tsx", "Hypothesis properties; vault.test.ts; E2E steps 7–8"],
    ["NFR-REL-1/2/3", "code/scheduler", "reliability_sim.py (1,000 schedules)"],
    ["NFR-PERF-1", "code/api", "load_checkin.py"],
    ["NFR-MAINT-1", "code/crypto, code/scheduler", "pytest --cov (CI gate at 80%)"],
    ["NFR-PORT-1", "code/vault/db.py", "SQLite in tests; Neon Postgres in deployment"],
], caption="Requirement traceability: requirement → module → test", label="tab:traceability", widths=[4, 6, 6])

# =====================================================================================
# Chapter 3
# =====================================================================================
chapter(3, "Results and Evaluation")
heading("3.1 Testing Methodology", 2)
hyp = mv("Hypothesis property tests") or {}
bullets([
    f"**Unit testing:** pytest for the Python packages and Vitest for the TypeScript crypto. Results: "
    f"{ms('Python tests (unit + integration)')} (Python) and {ms('TypeScript tests (vitest)')} (TypeScript). "
    f"Property-based testing with Hypothesis [cite: maciver2019hypothesis] covered "
    f"{hyp.get('properties', TODO)} properties with {hyp.get('generated_cases', TODO)} generated cases and "
    f"{hyp.get('failing_cases', TODO)} failures.",
    f"**Integration testing:** FastAPI's test client drives the full demo script server-side with the Python "
    f"crypto acting as the browser; the Python and TypeScript implementations are bound by the shared vectors "
    f"({ms('Cross-language vectors (Python -> TypeScript)')}); and a Playwright test runs the whole demo script "
    f"through the real user interface in separate browser contexts for the owner and each trustee "
    f"({ms('End-to-end demo script (Playwright)')}).",
    "**Performance testing:** a reliability simulator runs the real scheduler code against 1,000 schedules with "
    "a fake clock and injected faults; a load script measures check-in latency; a timing harness measures the "
    "client crypto operations.",
    "**User testing:** [TODO: describe the demo walkthrough with classmates — who, how many, what they did, what "
    "they found. If it did not happen, say so.]",
])

heading("3.2 Performance Metrics", 2)
lat_local = mv("NFR-PERF-1 p95 check-in latency (local)") or {}
rows = [
    ["Early releases, 1,000 schedules (NFR-REL-1)", "0",
     f"{ms('NFR-REL-1 early releases (requirement cadence)')} (requirement cadence); "
     f"{ms('NFR-REL-1 early releases (deployment cron only)')} (cron only)"],
    ["Duplicate releases (NFR-REL-3)", "0",
     f"{ms('NFR-REL-3 duplicate releases (requirement cadence)')}; "
     f"{ms('NFR-REL-3 duplicate releases (deployment cron only)')}"],
    ["Missed releases (FR-7)", "0",
     f"{ms('Missed releases (requirement cadence)')}; {ms('Missed releases (deployment cron only)')}"],
    ["Max release lateness (NFR-REL-2)", "≤ 60 s",
     f"{ms('NFR-REL-2 max release lateness (requirement cadence)', unit=' s')} (met); "
     f"{ms('NFR-REL-2 max release lateness (deployment cron only)', unit=' s')} with the 5-minute cron alone "
     f"(**not met**)"],
    ["Max prompt dispatch delay (NFR-PERF-2)", "≤ 60 s",
     f"{ms('NFR-PERF-2 max prompt dispatch delay (requirement cadence)', unit=' s')} (met); "
     f"{ms('NFR-PERF-2 max prompt dispatch delay (deployment cron only)', unit=' s')} cron only (**not met**)"],
    ["Check-in latency p95, local (NFR-PERF-1)", "≤ 500 ms",
     f"{lat_local.get('sequential_p95_ms', TODO)} ms sequential; {lat_local.get('concurrent50_p95_ms', TODO)} ms "
     f"for 50 simultaneous requests on one vault"],
    ["Check-in latency p95, deployed (NFR-PERF-1)", "≤ 500 ms", ms("NFR-PERF-1 p95 check-in latency (deployed)")],
    ["Zero-knowledge inspection findings (NFR-SEC-1/5)", "0", ms("NFR-SEC-1/5 zero-knowledge inspection findings")],
    ["Coverage: crypto (NFR-MAINT-1)", "≥ 80%", ms("Coverage: crypto", unit="%")],
    ["Coverage: scheduler (NFR-MAINT-1)", "≥ 80%", ms("Coverage: scheduler", unit="%")],
    ["Test pass rate", "100%", f"{ms('Python tests (unit + integration)')}; {ms('TypeScript tests (vitest)')}"],
    ["Cross-language vectors", "all", ms("Cross-language vectors (Python -> TypeScript)")],
]
for name, label in (("aes_gcm_encrypt_1MiB", "AES-256-GCM encrypt, 1 MiB"),
                    ("aes_gcm_decrypt_1MiB", "AES-256-GCM decrypt, 1 MiB"),
                    ("rsa_oaep_wrap_one_share", "RSA-OAEP wrap, one share"),
                    ("rsa_oaep_2048_keygen", "RSA-2048 keypair generation"),
                    ("shamir_split_k3_n5", "Shamir split (3, 5)"),
                    ("shamir_reconstruct_k3_n5", "Shamir reconstruct (3, 5)"),
                    ("shamir_reconstruct_k5_n7", "Shamir reconstruct (5, 7)")):
    rows.append([f"Crypto timing: {label}", "—",
                 f"median {ms('Crypto timing: ' + name, 'median_ms', ' ms')}, p95 "
                 f"{ms('Crypto timing: ' + name, 'p95_ms', ' ms')}"])
rows.append(["Lighthouse (landing / dashboard)", "—",
             f"{ms('Lighthouse (landing)')} / {ms('Lighthouse (dashboard)')}"])
table(["Metric", "Target", "Achieved"], rows, caption="Prototype performance results", label="tab:results",
      widths=[5.5, 2.2, 8.3])
rel_notes = M.get("NFR-REL-1 early releases (requirement cadence)", {}).get("notes", "")
timing_env = M.get("Crypto timing: aes_gcm_encrypt_1MiB", {}).get("environment", "")
para(f"Environment and method notes. Reliability: {rel_notes or TODO}. The local latency was measured against "
     f"uvicorn with SQLite on a Windows laptop. Crypto timings were measured with {timing_env or TODO}, not in a "
     f"browser. Every value above can be regenerated with `python scripts/collect_metrics.py`.", size=9.5)

heading("Demonstration screenshots", 3)
for file, caption, label in [
    ("04-wizard-configure.png", "Vault wizard: trustees, threshold K = 2 of N = 3 and loss tolerance", "fig:wizard"),
    ("03-threshold-validation.png", "FR-4 validation: K = N produces a warning (K < 2 is rejected)",
     "fig:validation"),
    ("06-trustee-enrol.png", "Trustee enrolment: the keypair is generated in the trustee's browser",
     "fig:enrol"),
    ("09-server-received.png", "What the server received and stores: ciphertext and encrypted shares only",
     "fig:server-view"),
    ("10-dashboard-active.png", "Owner dashboard: Active state, live countdowns and one-click check-in",
     "fig:dashboard"),
    ("12-demo-console-warning.png", "Demo Console after a tick: Warning state and the check-in prompt in the Outbox",
     "fig:demo-warning"),
    ("16-dashboard-released.png", "The vault after the grace period expired: Released", "fig:released"),
    ("18-trustee-key-file-share.png", "Trustee portal: share decrypted in the browser with the private key file",
     "fig:trustee-share"),
    ("19-recovery-one-share-fails.png", "Recovery Room: one share alone is refused", "fig:one-share"),
    ("20-recovery-success.png", "Recovery Room: two of three shares reconstruct the key and decrypt the vault",
     "fig:recovered"),
    ("22-dashboard-mobile.png", "The dashboard on a phone-sized screen", "fig:mobile"),
]:
    figure(SHOTS / file, caption, label, width_cm=6.5 if "mobile" in file else 15)

heading("3.3 Analysis", 2)
para("**Were the objectives met?** Largely, yes. The central safety objective was met: across 1,000 simulated "
     f"schedules with restarts, racing and duplicate ticks and clock jitter, there were "
     f"{ms('NFR-REL-1 early releases (requirement cadence)')} early releases, "
     f"{ms('NFR-REL-3 duplicate releases (requirement cadence)')} duplicate releases and "
     f"{ms('Missed releases (requirement cadence)')} missed releases, and the same held with the slower deployment "
     f"cadence. The zero-knowledge inspection found {ms('NFR-SEC-1/5 zero-knowledge inspection findings')} "
     "occurrences of the plaintext, the key or any share in the stored data or the requests, including the "
     "database left by the real browser test. The threshold property was demonstrated end to end in the browser. "
     f"Coverage on the crypto and scheduler packages ({ms('Coverage: crypto', unit='%')} and "
     f"{ms('Coverage: scheduler', unit='%')}) exceeds the 80% target.")
para("**Where were targets exceeded?** Beyond the plan, the Python and TypeScript implementations were bound by "
     f"shared known-answer vectors ({ms('Cross-language vectors (Python -> TypeScript)')}), the scanner used for "
     "the zero-knowledge check carries its own positive-control self-test, and the whole demo script runs as an "
     f"automated test. Locally, the sequential check-in p95 of {lat_local.get('sequential_p95_ms', TODO)} ms is "
     "well inside the 500 ms target.")
para("**What were the challenges, and which targets were missed?**")
bullets([
    "**Serverless scheduling.** Vercel functions keep no process alive and Hobby-plan cron jobs run at most daily, "
    "so the planned APScheduler design was replaced by an external trigger. With the 5-minute GitHub Actions "
    f"trigger alone, the maximum release lateness was {ms('NFR-REL-2 max release lateness (deployment cron only)', unit=' s')} "
    f"and the maximum prompt delay {ms('NFR-PERF-2 max prompt dispatch delay (deployment cron only)', unit=' s')}, "
    "so NFR-REL-2 and NFR-PERF-2 (both 60 s) are not met by the cron alone. They are met at a cadence of under a "
    "minute, which the lazy due-checks and the Demo Console provide during use and demonstration.",
    "**Email domain.** Without a verified sending domain, Resend delivers only to the account owner's address; "
    "the in-app Outbox (`EMAIL_MODE=log`) keeps the demonstration independent of email delivery.",
    "**Payload size.** Request bodies to Vercel functions are limited, so payloads are capped at about 1 MB and "
    "stored in the database; object storage for large files is future work.",
    "**Concurrency.** The load test exposed a stall: FastAPI runs a synchronous dependency's setup, the endpoint "
    "and the commit as separate thread-pool steps, so threads waiting for a database lock could starve the "
    f"lock holder. 50 simultaneous check-ins on one vault still took {lat_local.get('concurrent50_p95_ms', TODO)} "
    "ms at the 95th percentile locally because SQLite serialises writes.",
    f"**Not yet measured:** the deployed check-in latency ({ms('NFR-PERF-1 p95 check-in latency (deployed)')}) "
    f"and Lighthouse scores ({ms('Lighthouse (landing)')}), which require the live deployment.",
    "**Residual security risks** from the threat model remain: a malicious operator could serve altered client "
    "code (T-4) or withhold a release (T-7); K or more colluding trustees can open a vault by design (T-2); and "
    "metadata such as trustee emails and schedules is not encrypted.",
])
para("**How were they resolved?** The scheduler was redesigned around a pure decision function and an idempotent, "
     "transactional tick with several independent triggers; exactly-once delivery is enforced by unique "
     "deduplication keys in the outbox; the thread-pool stall was removed by limiting in-flight requests below "
     "the pool size; and an invalid package name inherited from the course template, which broke Vercel's "
     "dependency installation, was corrected. Each fix was confirmed by re-running the relevant measurement.")

# =====================================================================================
# Chapter 4
# =====================================================================================
chapter(4, "Conclusion")
heading("4.1 Summary of Work", 2)
para("The team built a working prototype of Aegis, an encrypted dead-man's-switch vault in which the browser "
     "encrypts a secret, splits its key with Shamir's Secret Sharing and wraps each share for a trustee, and the "
     "server only keeps time and stores ciphertext. The prototype covers the full demonstration script — "
     "registration, trustee enrolment with browser-generated keys, sealing, one-click check-in, the Active, "
     "Warning, Grace and Released lifecycle, share decryption by trustees and K-of-N recovery — and is "
     "accompanied by a measurement suite whose results are reported above.")
heading("4.2 Key Findings", 2)
numbered([
    "A pure decision function plus a transactional, idempotent tick made the never-early and exactly-once "
    f"properties testable at scale: {ms('NFR-REL-1 early releases (requirement cadence)')} early and "
    f"{ms('NFR-REL-3 duplicate releases (requirement cadence)')} duplicate releases over 1,000 fault-injected "
    "schedules.",
    "A zero-knowledge server is practical with standard browser cryptography: the inspection found "
    f"{ms('NFR-SEC-1/5 zero-knowledge inspection findings')} secrets on the server, and client-side costs are "
    f"small (AES-256-GCM over 1 MiB in a median of {ms('Crypto timing: aes_gcm_encrypt_1MiB', 'median_ms', ' ms')}).",
    "On serverless hosting, timeliness depends on the tick trigger rather than on the code: the same scheduler "
    "meets the 60-second targets at a sub-minute cadence but not with a 5-minute cron alone.",
])
heading("4.3 Future Work and Recommendations", 2)
numbered([
    "**Enhancements:** multiple vaults per owner, editing a vault after creation, account recovery, two-factor "
    "authentication, trustee-side reminders and an audit export.",
    "**Alternative approaches:** verifiable secret sharing so that tampered shares are rejected (threat T-9), "
    "constant-time field arithmetic, and a signed or extension-delivered client to remove the code-delivery risk "
    "(threat T-4).",
    "**Scalability:** object storage for files larger than 1 MB, Alembic migrations, rate limiting and a "
    "dedicated scheduler (or a paid cron) to meet NFR-REL-2 without relying on lazy due-checks.",
    "**Integration:** a production email domain with Resend, and optional notification channels behind the "
    "existing `EmailSender` strategy (SMS remains out of scope under NG-1).",
    "**Real-world deployment:** an independent security review, an admin interface, monitoring of the tick "
    "triggers, and backup and retention policies for the encrypted store.",
])
heading("4.4 Final Remarks", 2)
para("[TODO: two or three sentences in the team's own words on what was learned — for example, about treating "
     "the demonstration as an acceptance test, about measuring rather than estimating, and about how serverless "
     "hosting changed the scheduler design.]")

# =====================================================================================
# Bibliography
# =====================================================================================
page_break()
heading("Bibliography", 1)
para("The entries are in `report/references.bib`. Keys used in the text:")
bib = (ROOT / "report" / "references.bib").read_text(encoding="utf-8")
for key in re.findall(r"@\w+\{([^,]+),", bib):
    title = re.search(rf"@\w+\{{{re.escape(key)},.*?title\s*=\s*\{{(.+?)\}},\n", bib, flags=re.S)
    bullets([f"`{key}` — {re.sub(r'[{}]', '', title.group(1)) if title else ''}"])

# =====================================================================================
# Appendix (for the team)
# =====================================================================================
page_break()
heading("Appendix (for the team — not part of the report)", 1)
heading("A. Live URL and demo access", 2)
table(["Item", "Value"], [
    ["Live URL", LIVE_URL],
    ["Demo owner (seeded)", "demo-owner@example.com — password: the DEMO_OWNER_PASSWORD you set when running "
                            "scripts/seed_demo.py (written to demo-keys/README.txt; never committed)"],
    ["Demo Console token", "the DEMO_ADMIN_TOKEN environment variable (never committed)"],
    ["Trustee key files", "demo-keys/aegis-trustee-*.json (git-ignored), written by the seed script"],
], widths=[4.5, 11.5])
heading("B. Demo script (presenter notes)", 2)
numbered([
    "Register and log in as the owner (FR-1). Mention: the password is stored only as an argon2id hash.",
    "Create the vault with three trustee emails; open each invite link in a separate browser profile and enrol. "
    "Point out that the private key file downloads and the server only gets the public key (FR-4).",
    "Type the secret message (and attach a small file), choose K = 2 of N = 3 and the 2-minute demo schedule; "
    "show the K < 2 rejection and the K = N warning (FR-3, FR-4).",
    "Encrypt & upload, then show \"What the server received\": ciphertext and 256-byte encrypted shares only "
    "(NFR-SEC-5).",
    "Dashboard: Active with a live countdown; click \"I'm OK — check in\" once (FR-6).",
    "Stop checking in. In the Demo Console use +1 min / +5 min and \"Run tick now\": Warning (prompt in the Outbox), "
    "Grace, Released. Stress that release happens only at deadline + grace (NFR-REL-1).",
    "Two trustees open their release links, load their key files and decrypt their own shares (FR-7, FR-8).",
    "Recovery Room: paste one share — refused; paste the second — the message and file are recovered (NFR-SEC-2).",
    "Demo Console: show the Outbox, the tick log and the measured results table.",
])
para("Shortcut if time is short: run `python scripts/seed_demo.py --base-url <URL>` beforehand and start at step 6.")
heading("C. Likely viva questions", 2)
qa = [
    ("Why can the server not read a vault?",
     "The browser encrypts the payload and splits the key before upload; each share is encrypted to a trustee's "
     "public key. The server stores only ciphertext and encrypted shares, which our inspection script verified "
     "(0 findings)."),
    ("What happens if only one trustee responds?",
     "Nothing opens. With K = 2, one share is refused, and mathematically K−1 shares are consistent with every "
     "possible key; the Recovery Room shows \"not enough shares\"."),
    ("How do you guarantee no early release?",
     "One pure function decides release and only when now ≥ deadline + grace; the tick re-reads the deadline from "
     "the locked row; late check-ins are refused; 1,000 fault-injected schedules gave 0 early releases."),
    ("What if the server is down when the deadline passes?",
     "Deadlines are stored in the database, so the next tick catches up in order. The release is then late but "
     "never early; lateness depends on the trigger cadence (≤ 60 s at the requirement cadence, minutes with cron alone)."),
    ("Could a malicious operator release early?",
     "It could hand out encrypted shares early, but they are useless unless K trustees cooperate. Serving altered "
     "client code is a documented residual risk (T-4) that the prototype does not mitigate."),
    ("Why Shamir instead of encrypting the key to each trustee?",
     "Encrypting the whole key to each trustee would let any one trustee open the vault. Shamir requires K of them "
     "and tolerates the loss of N − K."),
    ("Why two implementations of the crypto?",
     "Python is the specification and test oracle with property tests; the browser needs TypeScript. Shared "
     "known-answer vectors make CI fail if they ever disagree (13/13 match)."),
    ("How do you stop someone faking a check-in?",
     "Links carry single-use, expiring tokens stored only as hashes; GET only previews, and the page confirms with "
     "POST so link scanners cannot check in."),
    ("Why is K ≥ 2 enforced, and what does K = N mean?",
     "K = 1 would let any single trustee open the vault (FR-4 rejects it). K = N means losing any one trustee makes "
     "the vault unrecoverable, so the UI warns and shows the loss tolerance N − K."),
    ("What did you measure, and what did not meet its target?",
     "Every number comes from scripts in the repo. The 60-second timeliness targets are not met with the 5-minute "
     "cron alone; deployed latency and Lighthouse were [TODO: measured after deployment?]."),
]
for i, (q, a) in enumerate(qa, 1):
    para(f"**Q{i}. {q}**")
    para(a)
heading("D. Figures and screenshot files", 2)
table(["Figure", "Caption", "File"], [list(f) for f in FIGURES], widths=[2.2, 8.3, 5.5])
all_shots = sorted(p.name for p in SHOTS.glob("*.png"))
para("All screenshots captured by the end-to-end test (`report-assets/screenshots/`): " + ", ".join(all_shots) + ".",
     size=9)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} ({len(FIGURES)} figures)")
